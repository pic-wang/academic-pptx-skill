"""House style helpers for academic .pptx decks (python-pptx).

Every visual rule in STYLE.md is encoded here. Build decks by calling these
helpers instead of setting fonts, colours and borders by hand.

Inline markup accepted by every text helper:
    **text**  -> red (C00000) bold run   (key phrase)
    *text*    -> red (C00000) regular run
    plain     -> 1A1A1A regular run

Canvas: 13.333 x 7.5 in (16:9). Font sizes and positions were measured from
the three reference decks (see STYLE.md).
"""
from __future__ import annotations

import copy
import json
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- constants
FONT = "Cambria"
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xC0, 0x00, 0x00)
GRAY = RGBColor(0xF2, 0xF2, 0xF2)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
AGENDA_INACTIVE = RGBColor(0xBF, 0xBF, 0xBF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
CONTENT_W = SLIDE_W - 2 * MARGIN      # 12.33 in
INDENT = 0.31                         # bullet indent (285750 EMU in the reference decks)

SIZE = {
    "title": 28,        # page title
    "section": 20,      # red-bullet section header
    "body": 20,         # body text
    "body_small": 18,   # side panels next to a figure, dense text
    "label": 14,        # labels inside diagrams
    "takehome": 20,
    "table": 18,
    "caption": 18,      # Fig. N. / Tab. N. captions and labels under images
    "cite": 10.5,
    "page": 14,
    "deck_title": 28,
    "authors": 20,
    "agenda": 24,
    "box_head": 18,     # concept box header
    "box_item": 16,     # concept box bullets
}

TITLE_Y = 0.5
TITLE_W = 9.85          # leaves the top-right corner free for the logo
SECTION_Y = 1.5         # first section header
BODY_Y = 2.0            # body under the first section header
TAKEHOME_BOTTOM = 6.85  # bottom edge of the take-home box
CITE_Y = 7.02
LOGO_W = 2.745          # institution logo width in the reference decks

ASSET_DIR = Path(__file__).resolve().parent / "assets"

_MARKUP = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def house_assets():
    """Local, untracked institution images listed in assets/assets.json.

    Returns {name: {"path", "x", "y", "w", ...}}; empty when the folder is
    absent. Typical names: polyu_logo (content-slide logo), title_corner,
    title_slogan, lab_logo (title-slide decorations).
    """
    meta = ASSET_DIR / "assets.json"
    if not meta.exists():
        return {}
    items = json.loads(meta.read_text(encoding="utf-8"))
    return {k: dict(v, path=str(ASSET_DIR / v["file"])) for k, v in items.items()
            if (ASSET_DIR / v["file"]).exists()}


# ---------------------------------------------------------------- deck setup
def new_presentation(template: str | Path | None = None) -> Presentation:
    """Blank 16:9 deck, or `template` with its existing slides removed."""
    if template and Path(template).exists():
        prs = Presentation(str(template))
        sldIdLst = prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
    else:
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W)
        prs.slide_height = Inches(SLIDE_H)
    return prs


def blank_slide(prs):
    """Add a slide on the layout with the fewest placeholders."""
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    return slide


# ---------------------------------------------------------------- text core
def _style_run(run, size, bold=False, color=TEXT, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # keep CJK / complex-script runs in Cambria too
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def add_runs(paragraph, text, size, bold=False, color=TEXT):
    """Write `text` into `paragraph`, honouring **red bold** / *red* markup."""
    for part in _MARKUP.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run()
            r.text = part[2:-2]
            _style_run(r, size, bold=True, color=RED)
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run()
            r.text = part[1:-1]
            _style_run(r, size, bold=bold, color=RED)
        else:
            r = paragraph.add_run()
            r.text = part
            _style_run(r, size, bold=bold, color=color)


def _plain(text):
    return _MARKUP.sub(lambda m: m.group(0).strip("*"), text)


def _text_width(text, size):
    """Upper estimate (in) of the rendered width of `text` in Cambria at `size` pt.

    0.0069 in per pt-character covers bold runs; regular text is ~8% narrower.
    """
    return len(_plain(text)) * size * 0.0069


def _textbox(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return tb, tf


def _set_bullet(paragraph, char="•", font="Arial", marL=INDENT, color=None):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(marL)))
    pPr.set("indent", str(-Inches(marL)))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if color is not None:
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(buClr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", font)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def _line_spacing(paragraph, pct):
    pPr = paragraph._p.get_or_add_pPr()
    ln = pPr.find(qn("a:lnSpc"))
    if ln is None:
        ln = etree.SubElement(pPr, qn("a:lnSpc"))
        pPr.insert(0, ln)
    for c in list(ln):
        ln.remove(c)
    etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(pct * 1000)))


# ---------------------------------------------------------------- page parts
def add_page_title(slide, text, x=MARGIN, y=TITLE_Y, w=TITLE_W):
    """28 pt Cambria bold black, top-left. No accent line underneath.

    Must fit on one line: about 54 characters at the default width.
    """
    plain = _plain(text)
    if len(plain) * SIZE["title"] * 0.0066 > w:
        import warnings
        warnings.warn(f"page title may wrap ({len(plain)} chars): {plain!r}")
    tb, tf = _textbox(slide, x, y, w, 0.6, anchor=MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, text, SIZE["title"], bold=True, color=BLACK)
    return tb


def add_section_header(slide, text, x=MARGIN, y=SECTION_Y, w=CONTENT_W):
    """Red bullet + 20 pt Cambria bold red text."""
    tb, tf = _textbox(slide, x, y, w, 0.45)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _set_bullet(p, color=RED)
    add_runs(p, text, SIZE["section"], bold=True, color=RED)
    return tb


def add_body(slide, items, x=MARGIN + INDENT, y=BODY_Y, w=CONTENT_W - INDENT, h=3.4,
             size=None, bullets=True, space_after=8):
    """Body paragraphs, left-aligned.

    `items`: list of str, or (level, str). Level 0 = black bullet, level 1 =
    en-dash sub-bullet. Pass bullets=False for plain lead-in sentences.
    At the default x, bullets sit under the section-header text.
    """
    size = size or SIZE["body"]
    tb, tf = _textbox(slide, x, y, w, h)
    first = True
    for it in items:
        level, text = it if isinstance(it, tuple) else (0, it)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        if bullets:
            _set_bullet(p, char="•" if level == 0 else "–",
                        marL=INDENT + 0.3 * level, color=TEXT)
        add_runs(p, text, size - (2 if level else 0))
    return tb


def add_takehome_box(slide, text, y=None, w=None, x=None, size=None):
    """Red dashed-border box at the bottom. Mandatory on every content slide.

    Width defaults to fit the text (single line) up to CONTENT_W; the box is
    centred horizontally and its bottom edge sits at TAKEHOME_BOTTOM unless
    `y` is given. Put the key phrase in **red bold** markup.
    """
    size = size or SIZE["takehome"]
    tw = _text_width(text, size)
    if w is None:
        w = min(CONTENT_W, 0.45 + tw)
    if x is None:
        x = (SLIDE_W - w) / 2
    lines = max(1, -(-tw // (w - 0.3)))
    h = 0.14 + lines * size * 1.2 / 72
    if y is None:
        y = TAKEHOME_BOTTOM - h
    tb, tf = _textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    tf.margin_left = tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, text, size)
    spPr = tb._element.spPr
    for el in spPr.findall(qn("a:ln")):
        spPr.remove(el)
    ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", "19050")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", "C00000")
    etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    return tb


def add_caption(slide, label, text, x, y, w, align=PP_ALIGN.CENTER):
    """18 pt caption: bold 'Fig. N.' / 'Tab. N.' prefix, single line.

    Pass label="" for a plain label under an image (e.g. "Air pollution").
    """
    tb, tf = _textbox(slide, x, y, w, 0.4)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    if label:
        r = p.add_run()
        r.text = label.rstrip() + " "
        _style_run(r, SIZE["caption"], bold=True)
    add_runs(p, text, SIZE["caption"])
    return tb


def add_citation(slide, text, x=0.1, y=CITE_Y, w=12.3):
    """Full reference in small muted text at the bottom-left."""
    tb, tf = _textbox(slide, x, y, w, 0.3)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, text, SIZE["cite"], color=MUTED)
    return tb


def add_page_number(slide, n):
    tb, tf = _textbox(slide, SLIDE_W - 0.8, SLIDE_H - 0.42, 0.7, 0.35)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_runs(p, str(n), SIZE["page"])
    return tb


def add_logo(slide, path, w=LOGO_W):
    """Institution logo in the top-right corner (optional)."""
    if path and Path(path).exists():
        return slide.shapes.add_picture(str(path), Inches(SLIDE_W - w - 0.15),
                                        Inches(0.10), width=Inches(w))
    return None


def add_figure(slide, path, x, y, w=None, h=None, caption=None):
    """Picture plus optional (label, text) caption directly beneath it.

    With both `w` and `h`, the image is scaled to fit inside that box with its
    aspect ratio kept, and centred horizontally in it. The caption height
    (0.4 in) is reserved inside the box.
    """
    if w and h:
        from PIL import Image
        with Image.open(path) as im:
            ar = im.width / im.height
        box_h = h - (0.4 if caption else 0)
        fw, fh = (w, w / ar) if w / ar <= box_h else (box_h * ar, box_h)
        x, w, h = x + (w - fw) / 2, fw, fh
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None,
                                   Inches(h) if h else None)
    if caption:
        label, text = caption
        add_caption(slide, label, text, x, y + Emu(pic.height).inches + 0.03,
                    Emu(pic.width).inches)
    return pic


def add_concept_boxes(slide, boxes, y=4.4, h=1.9, x=MARGIN, w=CONTENT_W, gap=0.35):
    """Row of rounded red-bordered boxes: [(header, [bullet, ...]), ...]."""
    n = len(boxes)
    bw = (w - gap * (n - 1)) / n
    shapes = []
    for i, (head, items) in enumerate(boxes):
        bx = x + i * (bw + gap)
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx),
                                    Inches(y), Inches(bw), Inches(h))
        sh.adjustments[0] = 0.12
        sh.fill.background()
        sh.line.color.rgb = RED
        sh.line.width = Pt(1.5)
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(6)
        add_runs(p, head, SIZE["box_head"], bold=True)
        for it in items:
            q = tf.add_paragraph()
            q.alignment = PP_ALIGN.LEFT
            q.space_after = Pt(4)
            _set_bullet(q, marL=0.25, color=TEXT)
            add_runs(q, it, SIZE["box_item"])
        shapes.append(sh)
    return shapes


# ---------------------------------------------------------------- tables
def _cell_border(tag, w_emu, color="000000"):
    ln = etree.Element(qn(tag))
    ln.set("w", str(w_emu))
    ln.set("cap", "flat")
    ln.set("cmpd", "sng")
    ln.set("algn", "ctr")
    if w_emu:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
    else:
        etree.SubElement(ln, qn("a:noFill"))
    return ln


def _apply_three_line_borders(table, rule_w=12700):
    """Black top rule, black header underline, black bottom rule. No grid.

    OOXML order trap: inside <a:tcPr>, lnL/lnR/lnT/lnB must precede the cell
    fill, or PowerPoint ignores the borders. This rebuilds tcPr in order.
    """
    n_rows = len(table.rows)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            fills = [copy.deepcopy(el) for el in tcPr
                     if el.tag in (qn("a:solidFill"), qn("a:noFill"))]
            for el in list(tcPr):
                tcPr.remove(el)
            top = rule_w if r in (0, 1) else 0
            bottom = rule_w if r in (0, n_rows - 1) else 0
            tcPr.append(_cell_border("a:lnL", 0))
            tcPr.append(_cell_border("a:lnR", 0))
            tcPr.append(_cell_border("a:lnT", top))
            tcPr.append(_cell_border("a:lnB", bottom))
            for el in fills:
                tcPr.append(el)


def add_table(slide, header, rows, x, y, w, col_widths=None, highlight=(),
              size=None, row_h=0.45):
    """Three-line academic table with zebra stripes.

    `highlight`: iterable of (row_idx, col_idx) data-cell positions (0-based,
    excluding the header) rendered red bold. Values may also use markup.
    Compact side tables: size=14, row_h=0.34.
    """
    size = size or SIZE["table"]
    n_rows, n_cols = len(rows) + 1, len(header)
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(row_h * n_rows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if attr in tblPr.attrib:
            del tblPr.attrib[attr]
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    widths = col_widths or [w / n_cols] * n_cols
    for r, values in enumerate([header] + list(rows)):
        for c, val in enumerate(values):
            if _text_width(str(val), size) + 0.16 > widths[c]:
                import warnings
                warnings.warn(f"table cell may wrap (row {r}, col {c}, {widths[c]:.2f} in): {val!r}")
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
    hl = set(highlight)
    for r in range(n_rows):
        values = header if r == 0 else rows[r - 1]
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            val = str(values[c])
            if r > 0 and (r - 1, c) in hl:
                val = f"**{val.strip('*')}**"
            add_runs(p, val, size, bold=(r == 0), color=BLACK if r == 0 else TEXT)
            if r == 0 or r % 2 == 0:
                cell.fill.background()
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = GRAY
    _apply_three_line_borders(tbl)
    return gf


# ---------------------------------------------------------------- whole slides
def title_slide(prs, title, authors, presenter=None, affiliation=None,
                venue=None, logo=None, decorations=()):
    """Centred title (key phrase via markup), authors, presenter block.

    `decorations`: images placed behind the text, each a dict with path, x,
    y, w (e.g. the title-page corner, slogan and lab logo from house_assets()).
    """
    s = blank_slide(prs)
    for d in decorations:
        s.shapes.add_picture(str(d["path"]), Inches(d["x"]), Inches(d["y"]), width=Inches(d["w"]))
    add_logo(s, logo)
    tb, tf = _textbox(s, 1.3, 2.6, 10.73, 1.6, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, title, SIZE["deck_title"], bold=True, color=BLACK)
    tb, tf = _textbox(s, 1.3, 4.6, 10.73, 0.45)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, authors, SIZE["authors"])
    if venue:
        tb, tf = _textbox(s, 1.3, 5.1, 10.73, 0.4)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, venue, SIZE["caption"], color=MUTED)
    lines = [l for l in (presenter and f"Presenter: {presenter}", affiliation) if l]
    for i, line in enumerate(lines):
        yy = 6.3 + i * 0.44
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.0), Inches(yy + 0.1),
                                 Inches(0.06), Inches(0.24))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RED
        bar.line.fill.background()
        bar.shadow.inherit = False
        tb, tf = _textbox(s, 8.12, yy, 5.1, 0.44, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        if line.startswith("Presenter: "):
            r = p.add_run()
            r.text = "Presenter: "
            _style_run(r, SIZE["authors"], bold=True, color=BLACK)
            add_runs(p, line[len("Presenter: "):], SIZE["authors"])
        else:
            add_runs(p, line, SIZE["authors"], bold=True, color=BLACK)
    return s


def agenda_slide(prs, sections, current, page=None, logo=None, title="Content"):
    """Section progress slide: current item black, others gray hollow bullets."""
    s = blank_slide(prs)
    add_logo(s, logo)
    add_page_title(s, title)
    tb, tf = _textbox(s, 1.9, 1.7, 9.7, 4.2, anchor=MSO_ANCHOR.MIDDLE)
    for i, sec in enumerate(sections):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        _line_spacing(p, 200)
        active = (i == current) or (sec == current)
        if active:
            _set_bullet(p, "•", "Arial", marL=0.4, color=BLACK)
        else:
            _set_bullet(p, "o", "Courier New", marL=0.4, color=BLACK)
        add_runs(p, sec, SIZE["agenda"], bold=True,
                 color=BLACK if active else AGENDA_INACTIVE)
    if page is not None:
        add_page_number(s, page)
    return s


def content_slide(prs, title, page=None, logo=None):
    """Blank content slide with page title, optional logo and page number."""
    s = blank_slide(prs)
    add_logo(s, logo)
    add_page_title(s, title)
    if page is not None:
        add_page_number(s, page)
    return s


def check_takehome(prs, exempt=(1,)):
    """Return 1-based indices of slides missing a dashed take-home box."""
    missing = []
    for i, s in enumerate(prs.slides, 1):
        if i in exempt:
            continue
        if not any(b'prstDash val="dash"' in etree.tostring(sh._element)
                   and sh.has_text_frame for sh in s.shapes):
            missing.append(i)
    return missing
